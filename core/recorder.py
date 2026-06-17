"""Signal recorder — accumulates timestamped samples and exports to Excel."""

import threading
from datetime import datetime
from typing import Optional

import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


class SignalRecorder:
    """Thread-safe recorder that appends (timestamp_s, voltage) pairs."""

    def __init__(self, sample_rate: float) -> None:
        """Initialize with a given sample rate (Hz)."""
        self._sample_rate = sample_rate
        self._lock = threading.Lock()
        self._timestamps: list[float] = []
        self._voltages: list[float] = []
        self._recording = False
        self._start_wall: Optional[datetime] = None
        self._elapsed_samples: int = 0

    # ------------------------------------------------------------------
    # Control
    # ------------------------------------------------------------------

    def start(self) -> None:
        """Begin recording; resets any previously accumulated data."""
        with self._lock:
            self._timestamps.clear()
            self._voltages.clear()
            self._elapsed_samples = 0
            self._start_wall = datetime.now()
            self._recording = True

    def stop(self) -> None:
        """Stop recording. Data remains accessible until next start()."""
        with self._lock:
            self._recording = False

    # ------------------------------------------------------------------
    # Data ingestion
    # ------------------------------------------------------------------

    def feed(self, chunk: np.ndarray) -> None:
        """Append a voltage chunk from the serial worker (called from QThread)."""
        with self._lock:
            if not self._recording:
                return
            n = len(chunk)
            dt = 1.0 / self._sample_rate if self._sample_rate > 0 else 1.0
            base = self._elapsed_samples * dt
            ts = [base + i * dt for i in range(n)]
            self._timestamps.extend(ts)
            self._voltages.extend(chunk.tolist())
            self._elapsed_samples += n

    def update_sample_rate(self, sample_rate: float) -> None:
        """Update sample rate mid-recording (takes effect on next chunk)."""
        with self._lock:
            self._sample_rate = sample_rate

    # ------------------------------------------------------------------
    # State
    # ------------------------------------------------------------------

    @property
    def is_recording(self) -> bool:
        """Return True if recording is active."""
        with self._lock:
            return self._recording

    @property
    def sample_count(self) -> int:
        """Return number of samples recorded so far."""
        with self._lock:
            return len(self._timestamps)

    @property
    def start_time(self) -> Optional[datetime]:
        """Return the wall-clock time when recording started."""
        with self._lock:
            return self._start_wall

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------

    def save_xlsx(self, path: str) -> int:
        """Write recorded data to an Excel file. Returns number of rows saved."""
        with self._lock:
            timestamps = list(self._timestamps)
            voltages = list(self._voltages)

        if not timestamps:
            raise ValueError("No data recorded yet.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Signal Recording"

        # Header style
        header_fill = PatternFill("solid", fgColor="1E1E1E")
        header_font = Font(bold=True, color="00BFFF", name="Consolas")
        center = Alignment(horizontal="center")

        headers = ["Timestamp (s)", "Voltage (V)"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

        # Data rows
        data_font = Font(name="Consolas", size=10)
        for row_idx, (t, v) in enumerate(zip(timestamps, voltages), start=2):
            ws.cell(row=row_idx, column=1, value=round(t, 8)).font = data_font
            ws.cell(row=row_idx, column=2, value=round(v, 6)).font = data_font

        # Column widths
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 16

        # Metadata sheet
        meta = wb.create_sheet("Metadata")
        meta_font = Font(name="Consolas", size=10)
        meta_rows = [
            ("Recording Start", self._start_wall.strftime("%Y-%m-%d %H:%M:%S") if self._start_wall else "—"),
            ("Sample Rate (Hz)", self._sample_rate),
            ("Total Samples", len(timestamps)),
            ("Duration (s)", round(timestamps[-1], 4) if timestamps else 0),
            ("Vmax (V)", round(max(voltages), 6)),
            ("Vmin (V)", round(min(voltages), 6)),
        ]
        for r, (k, v) in enumerate(meta_rows, start=1):
            meta.cell(row=r, column=1, value=k).font = Font(bold=True, name="Consolas", color="00BFFF")
            meta.cell(row=r, column=2, value=v).font = meta_font
        meta.column_dimensions["A"].width = 22
        meta.column_dimensions["B"].width = 24

        wb.save(path)
        return len(timestamps)
